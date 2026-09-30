#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
PRAXIS
Sistema inteligente de asignación de prácticas

Desarrollado por:
Diego J. Maldonado Guzmán
Profesor del Área de Derecho Penal de la Universidad de Málaga
Investigador del Instituto Andaluz Interuniversitario de Criminología - Sección Málaga
"""

import re
import unicodedata
from pathlib import Path
from io import BytesIO

import pandas as pd
import streamlit as st


# =========================================================
# CONFIGURACIÓN
# =========================================================

APP_NAME = "PRAXIS"

COL = {
    # Alumnos
    "id": "idalumno",
    "apellido": "apellido",
    "nombre": "nombre",
    "perfil": "perfil",
    "nota": "notamedia",

    # Plazas
    "pid": "idplaza",
    "pname": "nombreplaza",
    "cm": "capacidadmanana",
    "ct": "capacidadtarde",
    "profiles": "perfilesadmitidos",
}


# =========================================================
# NORMALIZACIÓN DE TEXTO
# =========================================================

def clean_col(x):
    """
    Normaliza los nombres de las columnas.

    Ejemplos:
    'Nota media' -> 'notamedia'
    'Preferencia 1' -> 'preferencia1'
    'Preferencia_1_turno' -> 'preferencia1turno'
    """
    return (
        str(x)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )


def normalizar_texto(x):
    """
    Normaliza texto para comparar nombres de plazas:
    - minúsculas
    - elimina tildes
    - elimina espacios repetidos
    """
    texto = str(x).strip().lower()

    texto = "".join(
        c
        for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )

    texto = " ".join(texto.split())

    return texto


# =========================================================
# LEER EXCEL
# =========================================================

def leer_excel(source):

    if source is None:
        raise Exception("No se seleccionó archivo.")

    try:

        if isinstance(source, (str, Path)):

            path = Path(source)

            if not path.exists():
                raise Exception(f"No existe:\n{path}")

            df = pd.read_excel(
                path,
                keep_default_na=False,
                engine="openpyxl"
            )

        else:

            if hasattr(source, "seek"):
                source.seek(0)

            df = pd.read_excel(
                source,
                keep_default_na=False,
                engine="openpyxl"
            )

    except Exception as e:
        raise Exception(f"No se pudo leer el archivo Excel:\n{e}")

    # Normalizar nombres de columnas
    df.columns = [clean_col(c) for c in df.columns]

    # Eliminar NaN
    df = df.where(pd.notnull(df), "")

    return df.astype(str)


# =========================================================
# UTILIDADES
# =========================================================

def to_float(x):

    try:
        return float(
            str(x)
            .strip()
            .replace(",", ".")
        )

    except (ValueError, TypeError):
        return 0.0


def to_int(x):

    try:
        return int(
            float(
                str(x)
                .strip()
                .replace(",", ".")
            )
        )

    except (ValueError, TypeError):
        return 0


def normalizar_perfil(p):

    p = normalizar_texto(p)

    mapa = {
        "educadora social": "educacion social",
        "educador social": "educacion social",
        "educacion social": "educacion social",

        "trabajadora social": "trabajo social",
        "trabajador social": "trabajo social",
        "trabajo social": "trabajo social",

        "criminologa": "criminologia",
        "criminologo": "criminologia",
        "criminologia": "criminologia",
    }

    return mapa.get(p, p)


def parse_profiles(x):
    """
    Si la celda está vacía o dice 'todos',
    la plaza admite cualquier perfil.
    """

    x = normalizar_texto(x)

    if x in ("", "todos", "todo", "all", "cualquiera"):
        return None

    # Permite:
    # trabajo social | educación social
    # trabajo social; educación social
    perfiles = re.split(r"[|;,]", x)

    return {
        normalizar_perfil(p)
        for p in perfiles
        if str(p).strip()
    }


def normalize_shift(x):

    x = normalizar_texto(x)

    if x in (
        "",
        "any",
        "cualquiera",
        "indiferente",
        "ambos"
    ):
        return "any"

    if x in (
        "m",
        "manana",
        "morning"
    ):
        return "morning"

    if x in (
        "t",
        "tarde",
        "afternoon"
    ):
        return "afternoon"

    # Si hay un valor desconocido, no restringimos el turno.
    return "any"


# =========================================================
# VALIDACIÓN DE COLUMNAS
# =========================================================

def validar_columnas(df, cols, nombre):

    faltan = [
        c for c in cols
        if c not in df.columns
    ]

    if faltan:

        raise Exception(
            f"Faltan columnas obligatorias en {nombre}:\n\n"
            + "\n".join(f"• {c}" for c in faltan)
        )


# =========================================================
# PREPARAR PLAZAS
# =========================================================

def preparar_plazas(df):

    validar_columnas(
        df,
        [
            COL["pid"],
            COL["pname"],
            COL["cm"],
            COL["ct"],
        ],
        "el archivo de PLAZAS"
    )

    plazas = {}
    alias = {}
    ocupacion = {}

    for _, r in df.iterrows():

        pid_original = str(r[COL["pid"]]).strip()

        if not pid_original:
            continue

        pid = normalizar_texto(pid_original)

        nombre_original = str(
            r[COL["pname"]]
        ).strip()

        nombre_normalizado = normalizar_texto(
            nombre_original
        )

        cm = to_int(r[COL["cm"]])
        ct = to_int(r[COL["ct"]])

        # Evitar capacidades negativas
        cm = max(cm, 0)
        ct = max(ct, 0)

        profiles = None

        if COL["profiles"] in df.columns:
            profiles = parse_profiles(
                r[COL["profiles"]]
            )

        plazas[pid] = {
            "id": pid,
            "id_original": pid_original,
            "nombre": nombre_original,
            "cm_total": cm,
            "ct_total": ct,
            "profiles": profiles,
        }

        ocupacion[pid] = {
            "cm": 0,
            "ct": 0,
        }

        # La preferencia puede contener el ID...
        alias[pid] = pid

        # ...o el nombre de la plaza.
        if nombre_normalizado:
            alias[nombre_normalizado] = pid

    if not plazas:
        raise Exception(
            "No se ha encontrado ninguna plaza válida "
            "en el archivo de plazas."
        )

    return plazas, alias, ocupacion


# =========================================================
# DETECTAR PREFERENCIAS
# =========================================================

def detectar_preferencias(df):
    """
    Detecta automáticamente:

    preferencia1
    preferencia2
    preferencia3
    ...
    preferencia20

    No existe límite de 3 preferencias.
    """

    prefs = []

    for columna in df.columns:

        c = clean_col(columna)

        m = re.fullmatch(
            r"preferencia([0-9]+)",
            c
        )

        if m:
            prefs.append(
                int(m.group(1))
            )

    return sorted(set(prefs))


# =========================================================
# ASIGNACIÓN
# =========================================================

def asignar(
    df,
    plazas,
    alias,
    ocupacion,
    prefer_morning=True
):

    # Perfil es opcional.
    columnas_obligatorias = [
        COL["id"],
        COL["apellido"],
        COL["nombre"],
        COL["nota"],
    ]

    validar_columnas(
        df,
        columnas_obligatorias,
        "el archivo de ALUMNOS"
    )

    # Si no existe perfil, lo creamos vacío.
    if COL["perfil"] not in df.columns:
        df[COL["perfil"]] = ""

    preferencias = detectar_preferencias(df)

    if not preferencias:
        raise Exception(
            "No se han encontrado preferencias.\n\n"
            "Las columnas deben llamarse, por ejemplo:\n"
            "Preferencia 1\n"
            "Preferencia 2\n"
            "Preferencia 3\n"
            "..."
        )

    # =====================================================
    # PRIORIDAD: NOTA DESCENDENTE
    # =====================================================

    df["_nota"] = df[
        COL["nota"]
    ].apply(to_float)

    # mergesort mantiene un orden estable.
    df = df.sort_values(
        by=[
            "_nota",
            COL["apellido"],
            COL["nombre"],
        ],
        ascending=[
            False,
            True,
            True,
        ],
        kind="mergesort"
    ).reset_index(drop=True)

    resultados = []

    # =====================================================
    # ALUMNO POR ALUMNO
    # =====================================================

    for _, alumno in df.iterrows():

        perfil = normalizar_perfil(
            alumno.get(COL["perfil"], "")
        )

        asignado = False
        motivo_final = "sin preferencia disponible"

        # ===============================================
        # PREFERENCIAS EN ORDEN: 1, 2, 3, 4...
        # ===============================================

        for i in preferencias:

            pref_col = f"preferencia{i}"
            turno_col = f"preferencia{i}turno"

            valor_preferencia = alumno.get(
                pref_col,
                ""
            )

            pref_normalizada = normalizar_texto(
                valor_preferencia
            )

            if not pref_normalizada:
                continue

            turno = normalize_shift(
                alumno.get(
                    turno_col,
                    "any"
                )
            )

            # ===========================================
            # LOCALIZAR PLAZA
            # ===========================================

            if pref_normalizada in alias:
                pid = alias[pref_normalizada]

            elif pref_normalizada in plazas:
                pid = pref_normalizada

            else:
                motivo_final = (
                    f"Preferencia {i}: plaza no encontrada "
                    f"({valor_preferencia})"
                )
                continue

            plaza = plazas[pid]

            # ===========================================
            # PERFIL
            # ===========================================

            perfiles_admitidos = plaza["profiles"]

            if (
                perfiles_admitidos is not None
                and perfil not in perfiles_admitidos
            ):
                motivo_final = (
                    f"Preferencia {i}: perfil incompatible"
                )
                continue

            # ===========================================
            # TURNOS
            # ===========================================

            if turno == "morning":

                orden_turnos = ["cm"]

            elif turno == "afternoon":

                orden_turnos = ["ct"]

            elif prefer_morning:

                orden_turnos = ["cm", "ct"]

            else:

                orden_turnos = ["ct", "cm"]

            turno_asignado = None

            # ===========================================
            # COMPROBAR CAPACIDAD
            # ===========================================

            for codigo_turno in orden_turnos:

                if codigo_turno == "cm":
                    capacidad = plaza["cm_total"]

                else:
                    capacidad = plaza["ct_total"]

                ocupadas = ocupacion[pid][
                    codigo_turno
                ]

                disponibles = (
                    capacidad - ocupadas
                )

                if disponibles > 0:

                    # MUY IMPORTANTE:
                    # ocupamos exactamente UNA plaza.
                    ocupacion[pid][
                        codigo_turno
                    ] += 1

                    asignado = True

                    turno_asignado = (
                        "Mañana"
                        if codigo_turno == "cm"
                        else "Tarde"
                    )

                    break

            # ===========================================
            # PLAZA CONSEGUIDA
            # ===========================================

            if asignado:

                resultados.append({
                    "ID": alumno[COL["id"]],
                    "Apellido": alumno[COL["apellido"]],
                    "Nombre": alumno[COL["nombre"]],
                    "Perfil": alumno.get(COL["perfil"], ""),
                    "Nota": alumno[COL["nota"]],
                    "Empresa": plaza["nombre"],
                    "Turno": turno_asignado,
                    "Preferencia": i,
                    "Estado": "asignado",
                })

                # IMPORTANTÍSIMO:
                # al conseguir una plaza dejamos de
                # revisar preferencias de este alumno.
                break

            motivo_final = (
                f"Preferencia {i}: sin plazas disponibles"
            )

        # ===============================================
        # SIN PLAZA
        # ===============================================

        if not asignado:

            resultados.append({
                "ID": alumno[COL["id"]],
                "Apellido": alumno[COL["apellido"]],
                "Nombre": alumno[COL["nombre"]],
                "Perfil": alumno.get(COL["perfil"], ""),
                "Nota": alumno[COL["nota"]],
                "Empresa": "",
                "Turno": "",
                "Preferencia": "",
                "Estado": motivo_final,
            })

    return pd.DataFrame(resultados)


# =========================================================
# GENERAR EXCEL
# =========================================================

def generar_excel(resultado):

    from openpyxl.styles import PatternFill, Font, Alignment

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        resultado.to_excel(
            writer,
            sheet_name="Asignaciones",
            index=False
        )

        ws = writer.book["Asignaciones"]

        # Colores
        fill_header = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        fill_ok = PatternFill(
            "solid",
            fgColor="C6EFCE"
        )

        fill_warn = PatternFill(
            "solid",
            fgColor="FFF2CC"
        )

        fill_bad = PatternFill(
            "solid",
            fgColor="F8D7DA"
        )

        # Cabecera
        for cell in ws[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = fill_header

            cell.alignment = Alignment(
                horizontal="center"
            )

        # Filas
        for row in ws.iter_rows(min_row=2):

            estado = row[-1].value
            pref = row[-2].value

            if (
                estado == "asignado"
                and pref == 1
            ):
                fill = fill_ok

            elif estado == "asignado":
                fill = fill_warn

            else:
                fill = fill_bad

            for cell in row:
                cell.fill = fill
                cell.font = Font(
                    color="000000"
                )

        # Anchura automática
        for cells in ws.columns:

            ancho = max(
                len(str(c.value or ""))
                for c in cells
            ) + 2

            ws.column_dimensions[
                cells[0].column_letter
            ].width = min(
                max(ancho, 10),
                45
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    return output.getvalue()


# =========================================================
# INTERFAZ STREAMLIT
# =========================================================

st.set_page_config(
    page_title="PRAXIS",
    page_icon="🎓",
    layout="wide"
)

st.title("🎓 PRAXIS")
st.subheader(
    "Sistema inteligente de asignación de prácticas"
)

st.caption(
    "Asignación por prioridad académica, preferencias, "
    "turnos y perfiles compatibles."
)


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:

    st.header("⚙️ Configuración")

    alumnos = st.file_uploader(
        "📄 Archivo de alumnos",
        type=["xlsx"],
        key="archivo_alumnos"
    )

    plazas_file = st.file_uploader(
        "🏢 Archivo de plazas",
        type=["xlsx"],
        key="archivo_plazas"
    )

    st.divider()

    prefer_morning = st.checkbox(
        "Preferir turno de mañana cuando sea indiferente",
        value=True
    )

    ejecutar = st.button(
        "▶ Ejecutar asignación",
        type="primary",
        use_container_width=True,
        disabled=not (
            alumnos is not None
            and plazas_file is not None
        ),
    )

    st.divider()

    st.caption(
        "Los alumnos se procesan de mayor a menor "
        "nota media."
    )


# =========================================================
# SESSION STATE
# =========================================================

if "resultado" not in st.session_state:
    st.session_state.resultado = None

if "excel" not in st.session_state:
    st.session_state.excel = None


# =========================================================
# EJECUTAR
# =========================================================

if ejecutar:

    try:

        with st.spinner(
            "Calculando asignaciones..."
        ):

            df_alumnos = leer_excel(
                alumnos
            )

            df_plazas = leer_excel(
                plazas_file
            )

            plazas, alias, ocupacion = (
                preparar_plazas(
                    df_plazas
                )
            )

            resultado = asignar(
                df_alumnos.copy(),
                plazas,
                alias,
                ocupacion,
                prefer_morning
            )

            excel = generar_excel(
                resultado
            )

            st.session_state.resultado = (
                resultado
            )

            st.session_state.excel = excel

        st.success(
            "✅ Asignación completada correctamente."
        )

    except Exception as e:

        st.session_state.resultado = None
        st.session_state.excel = None

        st.error(
            f"❌ Error durante la asignación:\n\n{e}"
        )


# =========================================================
# RESULTADOS
# =========================================================

resultado = st.session_state.resultado

if resultado is not None:

    total = len(resultado)

    asignados = int(
        (
            resultado["Estado"]
            == "asignado"
        ).sum()
    )

    primera = int(
        (
            (
                resultado["Estado"]
                == "asignado"
            )
            &
            (
                pd.to_numeric(
                    resultado["Preferencia"],
                    errors="coerce"
                ) == 1
            )
        ).sum()
    )

    sin_plaza = total - asignados

    # =====================================================
    # MÉTRICAS
    # =====================================================

    st.subheader("📊 Resumen de la asignación")

    a, b, c, d = st.columns(4)

    a.metric(
        "👥 Alumnos",
        total
    )

    b.metric(
        "✅ Asignados",
        asignados
    )

    c.metric(
        "⚠️ Sin plaza",
        sin_plaza
    )

    d.metric(
        "🥇 1.ª preferencia",
        primera
    )

    # =====================================================
    # TABLA
    # =====================================================

    st.subheader("📋 Resultados")

    def colorear(row):

        if (
            row["Estado"] == "asignado"
            and str(
                row["Preferencia"]
            ) in ("1", "1.0")
        ):

            estilo = (
                "background-color: #d4f4dd; "
                "color: #000000;"
            )

        elif row["Estado"] == "asignado":

            estilo = (
                "background-color: #fff3cd; "
                "color: #000000;"
            )

        else:

            estilo = (
                "background-color: #f8d7da; "
                "color: #000000;"
            )

        return [
            estilo
        ] * len(row)

    tabla_estilizada = (
        resultado.style.apply(
            colorear,
            axis=1
        )
    )

    st.dataframe(
        tabla_estilizada,
        use_container_width=True,
        hide_index=True,
        height=520
    )

    st.caption(
        "🟢 Verde: primera preferencia · "
        "🟡 Amarillo: otra preferencia · "
        "🔴 Rojo: sin plaza"
    )

    # =====================================================
    # DESCARGAR EXCEL
    # =====================================================

    st.download_button(
        "📥 Descargar asignaciones.xlsx",
        data=st.session_state.excel,
        file_name="asignaciones.xlsx",
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        type="primary",
        use_container_width=False
    )


# =========================================================
# CRÉDITOS
# =========================================================

st.markdown(
    """
    <div style="
        margin-top: 45px;
        padding: 20px;
        border-top: 1px solid rgba(128,128,128,0.35);
        text-align: center;
        font-size: 14px;
        line-height: 1.7;
    ">
        <strong>Desarrollado por:</strong><br>
        <strong>Diego J. Maldonado Guzmán</strong><br>
        Profesor del Área de Derecho Penal de la Universidad de Málaga<br>
        Investigador del Instituto Andaluz Interuniversitario de Criminología – Sección Málaga
    </div>
    """,
    unsafe_allow_html=True
)
